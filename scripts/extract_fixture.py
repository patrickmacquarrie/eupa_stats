"""Pull everything the salary engine needs (and the sheet's own answers) out of a
master-sheet .xlsx into one JSON fixture. Throwaway tooling for validation only."""
import json, re, sys, datetime, warnings
import openpyxl
from openpyxl.worksheet.formula import ArrayFormula
warnings.filterwarnings("ignore")

src, out, event_sheets = sys.argv[1], sys.argv[2], sys.argv[3].split("|")
PL = len(sys.argv) > 4 and sys.argv[4] == "pl"  # Premier League layout
wb = openpyxl.load_workbook(src)                    # formulas
wv = openpyxl.load_workbook(src, data_only=True)    # cached values

def f(c):
    x = c.value
    return x.text if isinstance(x, ArrayFormula) else x
def d(x):
    if isinstance(x, datetime.datetime): return x.date().isoformat()
    if isinstance(x, datetime.date): return x.isoformat()
    return x
def num(x):
    try: return float(x)
    except (TypeError, ValueError): return 0.0

setup, setupf = wv["Setup"], wb["Setup"]
se = wv["4. Stats Entry"]
agg, aggf = wv["5. Stats Aggregate"], wb["5. Stats Aggregate"]

# --- settings -------------------------------------------------------------
setting = {}
for r in setup.iter_rows(min_row=3, max_row=40, min_col=13 if PL else 14, max_col=14 if PL else 15, values_only=True):
    if isinstance(r[0], str): setting[r[0].strip()] = r[1]
w = [num(se.cell(4, c).value) for c in range(10, 18)]  # J..Q multipliers actually used
cap_extra = {}
AGG0 = 43 if PL else 42  # first weekly column (AQ in PL, AP otherwise)
for c in range(AGG0, AGG0 + 16):  # weekly cap formulas
    wk = agg.cell(6, c).value
    m = re.findall(r"\+(\d+)\s*$", str(f(aggf.cell(5, c)) or ""))
    if wk and m and "$15+" in str(f(aggf.cell(5, c))): cap_extra[int(wk)] = float(m[-1])
league = {
    "weights": dict(zip(["win", "goal", "assist", "secondAssist", "block", "drop", "throwaway", "gso"], w)),
    "tieWeightFactor": 0.5,
    "absence": {"pctOfInitialPerWeek": num(se["G4"].value), "pctRuleWeeks": 2,
                "thereafter": "avgGrowthPerGamePlayed", "floorAtSubGrowth": True},
    "matchesPerWeek": num(se["H4"].value),
    "capBuffer": num(setting.get("Initial Salary Cap Buffer")),
    "capExtraByWeek": cap_extra,
    "teamsForCapAverage": num(setting.get("Teams")),
}
if setting.get("Team Win Bonus / Team Win"):
    league["teamResultBonus"] = {"win": num(setting["Team Win Bonus / Team Win"]), "tieFactor": 0.5}

# --- teams & players --------------------------------------------------------
teams = []
for r in setup.iter_rows(min_row=3, max_row=9, max_col=5, values_only=True):
    if r[1]: teams.append({"name": r[1], "gm": r[4], "isSubTeam": r[4] == "Sub"})
gm_to_team = {t["gm"]: t["name"] for t in teams}
players = []
for r in setup.iter_rows(min_row=3, max_row=400, min_col=8 if PL else 9, max_col=11 if PL else 12, values_only=True):
    gender, name, sal, gm = r
    if PL: gender = "SubX" if str(gm).upper() == "SUB" else "X"   # open league: everyone one group
    if not name: continue
    players.append({"name": str(name).strip(), "gender": (gender or "").strip(),
                    "initialSalary": num(sal), "team": gm_to_team.get(gm), "isSub": str(gm).upper() == "SUB"})

# --- trades -----------------------------------------------------------------
trades = []
tr = wv["Trades"]
for r in tr.iter_rows(min_row=4, max_row=400, max_col=7, values_only=True):
    if r[0] is None or not r[2]: continue
    trades.append({"id": r[0], "afterWeek": r[1], "player": r[2], "fromGm": r[3], "toGm": r[4],
                   "capSpaceTrade": num(r[5]), "playerAdd": num(r[6])})

# --- schedule ---------------------------------------------------------------
schedule = []
for r in wv["Schedule " if "Schedule " in wv.sheetnames else "Schedule"].iter_rows(min_row=7, max_row=40, max_col=2, values_only=True):
    if isinstance(r[0], (int, float)) and float(r[0]).is_integer() and r[1]: schedule.append({"week": int(r[0]), "date": d(r[1])})

# --- raw events -------------------------------------------------------------
events = []
for sheet in event_sheets:
    for r in wv[sheet].iter_rows(min_row=2, values_only=True):
        if not r[0] or r[0] == "date": continue
        events.append({"date": d(r[0]), "clock": r[1], "statTeam": r[3], "otherTeam": r[4],
                       "statScore": num(r[5]), "otherScore": num(r[6]), "action": r[7],
                       "player": r[8], "lastPlayer": r[9], "secLastPlayer": r[10]})

# --- the sheet's own per-game rows (ground truth + human sub decisions) ------
cols = ["week", "game", "gameId", "team", "opp", "player", "subFlag", "subbedFor", "played", "win",
        "goals", "assists", "secondAssists", "blocks", "drops", "throwaways", "gso", "touches"]
entries = []
for r in se.iter_rows(min_row=8, values_only=True):
    if not r[5]: continue
    e = dict(zip(cols, r[:18]))
    e["growth"] = num(r[23])
    e["win"] = None if e["win"] is None else num(e["win"])
    for k in ["goals", "assists", "secondAssists", "blocks", "drops", "throwaways", "gso", "touches", "played"]:
        e[k] = num(e[k])
    entries.append(e)

# --- the sheet's weekly salaries --------------------------------------------
current_week = int(num(wv["Summary"]["A1"].value))
expected = {}
for r in agg.iter_rows(min_row=8, max_row=200, min_col=39 if PL else 38, max_col=58 if PL else 57, values_only=True):
    if not r[0] or r[0] in ("NA",): continue
    expected[str(r[0]).strip()] = {"before": r[2], "avgGrowth": r[3], "weeks": list(r[4:])}
caps = [agg.cell(5, c).value for c in range(AGG0, AGG0 + 16)]

json.dump({"source": src.split("/")[-1], "currentWeek": current_week, "league": league, "teams": teams,
           "players": players, "trades": trades, "schedule": schedule, "events": events,
           "sheetEntries": entries, "expected": {"salaries": expected, "capByWeek": caps}},
          open(out, "w"), indent=1, default=str)
print(out, "events", len(events), "entries", len(entries), "players", len(players),
      "trades", len(trades), "capExtra", cap_extra, "currentWeek", current_week)
